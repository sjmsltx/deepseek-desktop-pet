# -*- coding: utf-8 -*-
"""官方技能包与办公文档写入测试（Batch 3-2）

覆盖：三个官方技能包（pdf_tools / image_batch / file_organize）的真实文件处理、
"先预览后执行"的安全路径、以及 office_skill 新增的 Word / PPT 写入。
"""
import json
import os
import sys

import pytest

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE not in sys.path:
    sys.path.insert(0, BASE)

import plugin_manager as pm                                      # noqa: E402
import skill_pack as spk                                         # noqa: E402
import office_skill as off                                       # noqa: E402

OFFICIAL = ('office_report', 'pdf_tools', 'image_batch', 'file_organize')


def _pack_module(name):
    m = pm.PluginManager(os.path.join(BASE, 'plugins'))
    assert name in m.plugins, '%s 没被加载（%s）' % (name, m.rejected.get(name, ''))
    return m.plugins[name]['module']


# ---------------------------------------------------------------- 通用：清单合规

@pytest.mark.parametrize('name', OFFICIAL)
def test_official_manifest_is_valid(name):
    path = os.path.join(BASE, 'plugins', name, spk.MANIFEST_NAME)
    with open(path, encoding='utf-8-sig') as f:
        meta = json.load(f)
    ok, msg, norm = spk.validate_manifest(meta)
    assert ok, '%s 清单不合规：%s' % (name, msg)
    assert norm.get('builtin') is True, '%s 应标记 builtin' % name
    assert norm.get('permissions'), '%s 应声明权限（让使用者看得见它要什么）' % name
    assert not (set(norm['permissions']) - set(spk.PERMISSIONS)), norm['permissions']


@pytest.mark.parametrize('name', OFFICIAL)
def test_official_entry_passes_scan(name):
    path = os.path.join(BASE, 'plugins', name, spk.MANIFEST_NAME)
    with open(path, encoding='utf-8-sig') as f:
        meta = json.load(f)
    with open(os.path.join(BASE, 'plugins', name, meta['entry']), encoding='utf-8') as f:
        src = f.read()
    ok, msg, _needed = spk.scan_code(src, (meta.get('permissions') or {}).keys())
    assert ok, '%s 入口代码没过扫描：%s' % (name, msg)


# ---------------------------------------------------------------- PDF 工具

def _mk_pdf(path, pages=2):
    from pypdf import PdfWriter
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=200, height=200)
    with open(path, 'wb') as f:
        w.write(f)
    return path


def test_pdf_info_merge_split(tmp_path):
    mod = _pack_module('pdf_tools')
    a = _mk_pdf(str(tmp_path / 'a.pdf'), 2)
    b = _mk_pdf(str(tmp_path / 'b.pdf'), 1)

    out = mod.pdf_tool({'action': 'info', 'path': a})
    assert '2 页' in out, out

    out = mod.pdf_tool({'action': 'merge', 'paths': json.dumps([a, b]), 'name': '合并测试'})
    assert '已合并 2 个文件' in out and '共 3 页' in out, out
    merged = os.path.join(BASE, '输出', '合并测试.pdf')
    assert os.path.isfile(merged)
    from pypdf import PdfReader
    assert len(PdfReader(merged).pages) == 3

    out = mod.pdf_tool({'action': 'split', 'path': a, 'pages': '1,2', 'name': '拆a'})
    assert '已拆分 2 页' in out, out
    os.remove(merged)
    for n in ('拆a_第01页.pdf', '拆a_第02页.pdf'):
        p = os.path.join(BASE, '输出', n)
        if os.path.isfile(p):
            os.remove(p)


def test_pdf_text_warns_on_no_text_layer(tmp_path):
    """没有文字层的 PDF（字体没带 ToUnicode）→ 应明确告知像乱码，而不是把乱码当正文交出去

    实战背景（2026-09-20）：用 PyMuPDF 内置中文字体造的 PDF 提取出来是乱码，
    真因是**数据源没有可读文字层**（同扫描件），所以工具要如实告知。
    """
    import pymupdf
    mod = _pack_module('pdf_tools')
    p = str(tmp_path / 'no_unicode.pdf')
    doc = pymupdf.open()
    page = doc.new_page(width=300, height=300)
    page.insert_text((40, 60), '这是没有 ToUnicode 的中文内容', fontname='china-s', fontsize=14)
    doc.save(p)
    doc.close()
    out = mod.pdf_tool({'action': 'text', 'path': p})
    assert '像乱码' in out, out[:200]

    # 对照组：嵌入真实 TTF（有 ToUnicode）不该误报
    good = str(tmp_path / 'good.pdf')
    doc = pymupdf.open()
    page = doc.new_page(width=300, height=300)
    page.insert_text((40, 60), '这一页中文应当能正确提取', fontname='hei',
                     fontfile=r'C:\Windows\Fonts\simhei.ttf', fontsize=14)
    doc.subset_fonts()
    doc.save(good)
    doc.close()
    out2 = mod.pdf_tool({'action': 'text', 'path': good})
    assert '像乱码' not in out2 and '这一页中文应当能正确提取' in out2, out2[:200]


def test_pdf_missing_file_and_bad_action(tmp_path):
    mod = _pack_module('pdf_tools')
    assert '找不到文件' in mod.pdf_tool({'action': 'info', 'path': str(tmp_path / 'none.pdf')})
    assert '未知 action' in mod.pdf_tool({'action': 'nope', 'path': _mk_pdf(str(tmp_path / 'c.pdf'))})


# ---------------------------------------------------------------- 图片批处理

def _mk_img(path, size=(100, 80), mode='RGB', fmt=None):
    from PIL import Image
    Image.new(mode, size, 'white').save(path, format=fmt)
    return path


def test_image_info_resize_convert_rename(tmp_path):
    mod = _pack_module('image_batch')
    src = tmp_path / 'imgs'
    src.mkdir()
    _mk_img(str(src / 'one.png'), (200, 100))
    _mk_img(str(src / 'two.png'), (100, 100))

    out = mod.image_batch({'action': 'info', 'src': str(src)})
    assert '共 2 张图片' in out and '200×100' in out, out

    out = mod.image_batch({'action': 'resize', 'src': str(src), 'width': 50, 'name': '缩放测试'})
    assert '已缩放 2 张' in out, out
    from PIL import Image
    with Image.open(os.path.join(BASE, '输出', '缩放测试', 'one.png')) as im:
        assert im.width == 50 and im.height == 25        # 等比

    out = mod.image_batch({'action': 'convert', 'src': str(src), 'format': 'jpg', 'name': '转jpg测试'})
    assert '已转换 2 张为 jpg' in out, out
    assert os.path.isfile(os.path.join(BASE, '输出', '转jpg测试', 'one.jpg'))

    out = mod.image_batch({'action': 'rename', 'src': str(src), 'prefix': '照片', 'name': '重命名测试'})
    assert '已复制并重命名 2 张' in out, out
    assert os.path.isfile(os.path.join(BASE, '输出', '重命名测试', '照片_001.png'))
    assert (src / 'one.png').exists(), '原图不该被改动'


def test_image_bad_src_and_format(tmp_path):
    mod = _pack_module('image_batch')
    assert '没找到图片' in mod.image_batch({'action': 'info', 'src': str(tmp_path / '空目录')})
    src = tmp_path / 'i'; src.mkdir(); _mk_img(str(src / 'x.png'))
    assert '不支持的格式' in mod.image_batch({'action': 'convert', 'src': str(src), 'format': 'gif2'})


# ---------------------------------------------------------------- 文件整理

def test_organize_preview_then_confirm(tmp_path):
    mod = _pack_module('file_organize')
    d = tmp_path / '乱'
    d.mkdir()
    for n in ('a.png', 'b.jpg', 'c.pdf', 'd.txt', 'e.zip'):
        (d / n).write_text('x', encoding='utf-8')

    out = mod.organize_files({'action': 'preview', 'src': str(d)})
    assert '整理计划' in out and '图片/' in out and '文档/' in out, out
    assert sorted(p.name for p in d.iterdir()) == ['a.png', 'b.jpg', 'c.pdf', 'd.txt', 'e.zip'], '预览不该动文件'

    out = mod.organize_files({'action': 'run', 'src': str(d)})      # 没带 confirm
    assert 'confirm=true' in out and (d / 'a.png').exists(), '没确认就不许动'

    out = mod.organize_files({'action': 'run', 'src': str(d), 'confirm': True})
    assert '已移动 5 个' in out, out
    assert (d / '图片' / 'a.png').exists() and (d / '文档' / 'c.pdf').exists()


def test_organize_dupes_and_rename(tmp_path):
    mod = _pack_module('file_organize')
    d = tmp_path / 'dup'
    d.mkdir()
    (d / 'x1.txt').write_text('same', encoding='utf-8')
    (d / 'x2.txt').write_text('same', encoding='utf-8')
    (d / 'y.txt').write_text('diff', encoding='utf-8')

    out = mod.organize_files({'action': 'dupes', 'src': str(d)})
    assert '发现 1 组重复' in out and '没有删除任何东西' in out, out
    assert (d / 'x1.txt').exists() and (d / 'x2.txt').exists(), '只报告不删除'

    out = mod.organize_files({'action': 'rename', 'src': str(d), 'prefix': '文'})
    assert 'confirm=true' in out, '重命名也要确认'
    out = mod.organize_files({'action': 'rename', 'src': str(d), 'prefix': '文', 'confirm': True})
    assert '已重命名 3 个' in out, out


def test_organize_by_date(tmp_path):
    mod = _pack_module('file_organize')
    d = tmp_path / 'bydate'
    d.mkdir()
    (d / 'f.txt').write_text('x', encoding='utf-8')
    out = mod.organize_files({'action': 'preview', 'src': str(d), 'by': 'date'})
    import time
    assert time.strftime('%Y-%m') in out, out


# ---------------------------------------------------------------- Word / PPT 写入

def test_make_doc_creates_readable_docx(tmp_path):
    out = str(tmp_path / '报告.docx')
    msg, ok = off.make_doc(out, '测试报告', ['第一段', {'style': 'Heading 1', 'text': '小节'},
                                            {'style': 'List Bullet', 'text': '要点'}],
                           ['列A', '列B'], [['1', '2']], subtitle='副标题')
    assert ok, msg
    from docx import Document
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert '测试报告' in texts and '第一段' in texts and '副标题' in texts
    assert '要点' in texts, '项目符号段落没写进去'
    assert len(doc.tables) == 1 and doc.tables[0].rows[0].cells[1].text == '列B'


def test_make_slides_creates_readable_pptx(tmp_path):
    out = str(tmp_path / '演示.pptx')
    msg, ok = off.make_slides(out, '演示标题', [
        {'title': '第一页', 'bullets': ['要点一', '要点二'], 'notes': '讲稿备注'},
        {'title': '第二页', 'bullets': ['结论']},
    ], subtitle='副标题')
    assert ok, msg
    from pptx import Presentation
    prs = Presentation(out)
    assert len(prs.slides._sldIdLst) == 3, '封面 + 2 页'
    assert prs.slides[0].shapes.title.text == '演示标题'
    assert '要点一' in prs.slides[1].placeholders[1].text
    assert prs.slides[1].notes_slide.notes_text_frame.text == '讲稿备注'


def test_make_doc_missing_lib_message(tmp_path, monkeypatch):
    import builtins
    real = builtins.__import__

    def fake(name, *a, **k):
        if name == 'docx':
            raise ImportError('no docx')
        return real(name, *a, **k)
    monkeypatch.setattr(builtins, '__import__', fake)
    msg, ok = off.make_doc(str(tmp_path / 'x.docx'), 'T', ['p'])
    assert not ok and 'python-docx' in msg
